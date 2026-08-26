#!/usr/bin/env python3
import argparse,csv,json
from pathlib import Path
from openpyxl import load_workbook

ACTION={'Transfer remaining order details':'transfer','Archive / do not transfer':'archive','Cancel / do not transfer':'cancel','Waiting on customer':'hold_customer','Waiting on Darren':'hold_warehouse'}
def main():
 p=argparse.ArgumentParser();p.add_argument('workbook');p.add_argument('orders_csv');p.add_argument('output');a=p.parse_args()
 wb=load_workbook(a.workbook,data_only=True,read_only=True);ws=wb['Open Orders Summary'];rows=list(ws.iter_rows(min_row=5,values_only=True))
 ids={}
 for r in csv.DictReader(open(a.orders_csv,encoding='utf-8-sig')):
  if r['Name'] and r['Id']: ids.setdefault(r['Name'],set()).add(r['Id'])
 decisions=[]
 for r in rows:
  order=str(r[0] or '');raw=str(r[15] or '');action=ACTION.get(raw,'undecided')
  order_ids=ids.get(order,set())
  if len(order_ids)!=1: raise SystemExit(f'{order}: expected one immutable ID, got {order_ids}')
  decisions.append({'order_number':order,'shopify_order_id':next(iter(order_ids)),'action':action})
 if len(decisions)!=27: raise SystemExit(f'expected 27 decisions, got {len(decisions)}')
 Path(a.output).write_text(json.dumps({'decisions':decisions},indent=2)+'\n')
 from collections import Counter
 print(json.dumps({'orders':len(decisions),'actions':Counter(x['action'] for x in decisions)},default=dict))
if __name__=='__main__':main()
