#!/usr/bin/env python3
import argparse, json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook


def records(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['Product Launch Decisions']
    headers = [c.value for c in ws[4]]
    idx = {h: i for i, h in enumerate(headers) if h}
    out = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[idx['Product Title']]:
            continue
        skus = [s.strip() for s in str(row[idx['SKUs']] or '').split(';') if s.strip()]
        out.append({
            'title': row[idx['Product Title']],
            'handle': row[idx['Handle']],
            'skus': skus,
            'decision': row[idx['Damon Final Decision']],
        })
    return out


def visibility(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['Product Visibility']
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h}
    return {
        row[idx['Product Title']]: row[idx['Damon Visibility Decision']]
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row[idx['Product Title']]
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument('product_workbook')
    p.add_argument('control_workbook')
    p.add_argument('output')
    args = p.parse_args()
    rows = records(args.product_workbook)
    vis = visibility(args.control_workbook)
    if len(rows) != 175:
        raise SystemExit(f'expected 175 decisions, got {len(rows)}')
    for row in rows:
        row['visibility'] = vis.get(row['title'])
        if 'NGPF7000' in row['skus']:
            row['shipping_weight_lb_by_sku'] = {'NGPF7000': 8}
    counts = Counter(r['decision'] for r in rows)
    if counts != Counter({'Launch': 140, 'Do Not Launch': 34, 'Archive': 1}):
        raise SystemExit(f'unexpected decision counts: {counts}')
    if any(not r['visibility'] for r in rows):
        raise SystemExit('missing visibility decisions')
    output = {'source': 'Damon handoff 2026-08-25', 'counts': dict(counts), 'products': rows}
    Path(args.output).write_text(json.dumps(output, indent=2) + '\n')
    print(json.dumps({'output': args.output, 'products': len(rows), 'counts': dict(counts)}))

if __name__ == '__main__':
    main()
