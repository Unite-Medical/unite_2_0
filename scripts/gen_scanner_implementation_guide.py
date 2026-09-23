#!/usr/bin/env python3
"""Generate the Unite WMS scanner buy and implementation guide on Desktop."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path.home() / "Desktop/Unite_WMS_Scanner_Buy_and_Implementation_Guide.docx"
NAVY = RGBColor(15, 42, 74)
TEAL = RGBColor(26, 110, 142)
GREEN = RGBColor(30, 125, 58)
AMBER = RGBColor(154, 106, 0)
RED = RGBColor(176, 42, 42)
GREY = RGBColor(80, 88, 96)
WHITE = RGBColor(255, 255, 255)
LIGHT = "EAF1F5"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_cell(cell, text, bold=False, color=None, size=8.8):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Aptos"
    if color:
        r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        style_cell(t.rows[0].cells[i], h, True, WHITE, 8.5)
        shade(t.rows[0].cells[i], "0F2A4A")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            style_cell(cells[i], value, bold=(i == 0), color=NAVY if i == 0 else None)
            if len(t.rows) % 2 == 0:
                shade(cells[i], "F5F8FA")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    return t


def title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = NAVY
    if subtitle:
        p2 = doc.add_paragraph()
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(11)
        r2.font.color.rgb = GREY


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(17)
    r.font.color.rgb = NAVY
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12.5)
    r.font.color.rgb = TEAL


def body(doc, text, bold=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(10)
    r.font.name = "Aptos"
    if color:
        r.font.color.rgb = color
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(9.7)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def callout(doc, heading, text, color="EAF1F5"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    shade(t.cell(0, 0), color)
    p = t.cell(0, 0).paragraphs[0]
    r = p.add_run(heading + ": ")
    r.bold = True
    r.font.color.rgb = NAVY
    p.add_run(text)
    margins(t.cell(0, 0), 120, 140, 120, 140)
    return t


d = Document()
for section in d.sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

normal = d.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)

# Cover
title(d, "Unite Medical WMS Scanner Buy and Implementation Guide",
      "The exact hardware, product controls, warehouse workflow, and launch tests for scan in, scan out, lot, expiration, backorders, and shipping blocks.")
callout(d, "Buy this now", "Two Zebra DS2278 cordless 2D scanner kits, one Zebra DS2208 corded 2D backup, one Zebra ZD421d 203 dpi direct thermal label printer, 4 x 2 inch lot labels, and 4 x 6 inch shipping labels.", "DDEFE5")
body(d, "This setup reads ordinary UPC barcodes plus GS1 DataMatrix and UDI codes that can carry GTIN, lot, serial, and expiration data. A 1D laser scanner is not sufficient for Unite's medical inventory requirements.", bold=True)

h1(d, "1. Purchase list")
rows = [
    ("Zebra DS2278 cordless 2D scanner kit", "Receiving and pallet scanning. Bluetooth HID, cradle, USB cable, standard-range 1D/2D imager.", "2", "$230 to $340", "$460 to $680"),
    ("Zebra DS2208 corded 2D scanner kit", "Backup and fixed packing station. USB keyboard-wedge mode, no charging dependency.", "1", "$110 to $190", "$110 to $190"),
    ("Zebra ZD421d, 203 dpi, direct thermal", "Print Unite lot labels and 4 x 6 shipping labels. Choose USB plus Ethernet model if warehouse network printing is desired.", "1", "$450 to $700", "$450 to $700"),
    ("4 x 2 direct thermal labels", "Lot, expiration, bin, and internal GS1 labels.", "6 rolls", "$10 to $25", "$60 to $150"),
    ("4 x 6 direct thermal labels", "Parcel and LTL shipping labels.", "6 rolls", "$15 to $30", "$90 to $180"),
    ("Spare DS2278 battery", "Avoids a dead scanner stopping receiving or picking.", "1", "$35 to $60", "$35 to $60"),
]
table(d, ["Item", "Purpose", "Qty", "Approx. each", "Approx. total"], rows, [1.8, 2.75, 0.55, 0.85, 0.9])
body(d, "Starter hardware range: $1,205 to $1,960. Formula: scanner kits + backup scanner + printer + both label stocks + spare battery. Confirm current distributor pricing before purchase.", bold=True)

h2(d, "Do not buy yet")
bullet(d, "Do not buy 1D-only laser scanners. They cannot read GS1 DataMatrix lot and expiration data.")
bullet(d, "Do not buy $1,000-plus rugged mobile computers for the first rollout. Add Zebra TC22 or TC27 devices only after fixed and Bluetooth stations prove the workflow and roaming volume justifies them.")
bullet(d, "Do not buy a proprietary inventory scanner that requires its own cloud inventory system. Unite WMS remains the system of record.")

d.add_section(WD_SECTION.NEW_PAGE)
section_2 = h1(d, "2. Why these scanners")
table(d, ["Capability", "DS2278 / DS2208", "Why Unite needs it"], [
    ("UPC and Code 128", "Yes", "Recognize ordinary products and internal labels."),
    ("GS1 DataMatrix / UDI", "Yes", "Capture GTIN, lot, expiration, and serial from medical packaging."),
    ("Bluetooth roaming", "DS2278", "Scan cartons and pallets away from the desk."),
    ("USB keyboard wedge", "Both", "Works in the browser without a proprietary driver."),
    ("Corded fallback", "DS2208", "Packing keeps working if wireless hardware is unavailable."),
])

h1(d, "3. Product setup controls")
body(d, "Damon must be able to configure tracking requirements for every product or variant. The system should not assume every item has the same regulatory or warehouse data.")
table(d, ["Product setting", "Values", "Operational effect"], [
    ("Lot tracking", "Not tracked / Optional / Required", "If Required, the item cannot ship without a lot assignment."),
    ("Expiration tracking", "Not tracked / Optional / Required", "If Required, the item cannot ship without an expiration date."),
    ("Serial / UDI tracking", "Not tracked / Optional / Required", "If Required, capture each serialized unit or UDI."),
    ("Minimum remaining shelf life", "Days", "Blocks allocation of inventory too close to expiration for the customer or product rule."),
    ("FEFO", "On / Off", "Suggests the first-expiring eligible lot during picking."),
    ("Photo evidence", "Optional / Required", "Requires packaging or damage photo at receiving/return."),
    ("Temperature control", "None / Refrigerated / Frozen / Custom", "Adds storage and shipping handling requirements."),
    ("Parcel profile", "Unit and case weight/dimensions", "Drives cartonization, parcel/pallet decision, and carrier rates."),
])
callout(d, "Hard shipping rule", "If a product requires lot, expiration, serial, or UDI and the required value is missing, the shipment cannot be completed. There is no warehouse override that bypasses the requirement. An authorized manager must correct the inventory record first.", "F9E3E1")

h1(d, "4. Receiving and backorder scan in")
for step in [
    "Open Receiving and select the purchase order, inbound shipment, transfer, or blind receipt.",
    "Scan the PO or inbound reference, then scan the carton or product barcode.",
    "Resolve GTIN to the Unite SKU. If the product is unknown, stop and route it to SKU resolution.",
    "Parse lot, expiration, and serial from GS1/UDI where present.",
    "Apply the product's required-field policy. Missing required values must be entered from the label, captured by photo/OCR and confirmed, or printed onto a Unite internal label before inventory becomes available.",
    "Enter or scan quantity, warehouse, bin, owner, and carton/pallet reference.",
    "Post receipt through the WMS ledger with an idempotency key. Never update on-hand directly.",
    "Reconcile ordered versus received quantities and flag short, over, damaged, and unexpected lines.",
    "If the received SKU has paid backorders, alert admin/assigned account reps and allocate according to backorder priority. Do not silently mark backorders shipped.",
    "Create each backorder suborder and invoice only when its allocated batch is ready for the next payment/shipment step.",
]:
    numbered(d, step)

h2(d, "Backorder receipt event")
body(d, "A receipt for a backordered SKU must emit one durable event. That event connects receiving, inventory, backorder allocation, admin notification, customer notification, suborder creation, invoicing, and later shipment. This closes the current gap where inventory can arrive without the business knowing which customer obligation it satisfies.")

h1(d, "5. Picking and scan out")
for step in [
    "Open the paid and allocated pick task.",
    "Scan the order or pick ticket.",
    "WMS displays the required SKU, quantity, bin, and FEFO lot.",
    "Scan the product and lot/UDI. Reject wrong SKU, wrong lot, expired lot, missing required field, or quantity above the reservation.",
    "Record picker, timestamp, station, lot genealogy, and quantity. Picking changes reservation state but does not yet remove on-hand inventory.",
    "At packing, re-scan products, record carton/pallet dimensions and weight, and generate the correct labels and paperwork.",
    "Only confirmed carrier handoff or shipment-manifest close posts the ship movement and reduces on-hand inventory.",
]:
    numbered(d, step)

h1(d, "6. Internal label fallback")
body(d, "When a manufacturer barcode does not carry required lot or expiration data, Unite should print an internal GS1-compatible label after staff confirms the values from the product packaging or OCR-assisted photo. The internal label should include human-readable SKU, lot, expiration, owner, and barcode content.")
bullet(d, "Preferred capture: manufacturer GS1 DataMatrix or UDI scan.")
bullet(d, "Second choice: Unite-generated internal GS1 label.")
bullet(d, "Third choice: photo plus OCR with staff confirmation.")
bullet(d, "Last resort: manual entry, visibly flagged in the audit record.")

h1(d, "7. Implementation work")
table(d, ["Area", "Required change", "Acceptance result"], [
    ("Product master", "Add configurable lot, expiration, serial/UDI, shelf-life, FEFO, temperature, and package-data policies.", "Damon can make each field optional or required per product."),
    ("Receiving", "Route all receipts through PO/inbound reconciliation and the append-only WMS ledger.", "One scan creates lot, movement, scan provenance, and reconciliation status."),
    ("Backorders", "Emit stock-arrived event, allocate paid obligations, alert admin/rep, and create batch suborders.", "Incoming backordered stock is never orphaned."),
    ("Picking", "Validate product, lot, expiry, quantity, and required fields against reservation.", "Invalid or incomplete items cannot be shipped."),
    ("Packing", "Re-scan, cartonize, collect actual dimensions/weight, and create parcel/LTL artifacts.", "Rates and labels use physical shipment truth."),
    ("Shipping", "Post stock movement only at carrier handoff/manifest close.", "Label creation never falsely means shipped."),
    ("Audit", "Record actor, device/station, raw barcode, parsed fields, manual corrections, and reason.", "Every inventory event is traceable."),
])

h1(d, "8. Installation and configuration")
for step in [
    "Unbox the DS2278 kits, connect each cradle by USB, and pair the scanner to its cradle.",
    "Configure USB HID keyboard mode, carriage return after scan, GS1 DataMatrix, Code 128, UPC-A/E, and transmission of FNC1/group separators.",
    "Pair one DS2278 to a tablet only if the receiving station is mobile. Keep the cradle connected for charging.",
    "Connect the DS2208 to the packing computer as the always-available wired fallback.",
    "Install the ZD421d, set 203 dpi direct thermal media, calibrate 4 x 2 and 4 x 6 label sizes, and print test labels.",
    "Create WMS device records for RECEIVING-1, PICK-1, and PACK-1. Record scanner serial number and station assignment.",
    "Run one known GS1 barcode through the browser and verify GTIN, lot, and expiration parse correctly.",
]:
    numbered(d, step)

h1(d, "9. Go-live test")
tests = [
    ("Receive known SKU with GS1 lot/expiration", "One scan fills SKU, lot, and expiration and posts one receipt movement."),
    ("Scan same receipt twice", "Idempotency prevents duplicate stock."),
    ("Receive item missing required lot", "Inventory remains blocked until lot is captured and confirmed."),
    ("Receive a paid backorder SKU", "Admin/rep alert appears and the correct obligation is identified."),
    ("Pick wrong SKU", "System blocks the pick."),
    ("Pick required-lot item without lot", "System blocks shipping."),
    ("Pick expired lot", "System blocks the pick and suggests an eligible lot."),
    ("Pack shipment", "Actual cartons, dimensions, and weight are recorded before label creation."),
    ("Create label", "Inventory is still on-hand until carrier handoff."),
    ("Confirm carrier handoff", "One ship movement posts, reservation clears, and lot genealogy records customer/order."),
    ("Recall test", "A lot search returns every affected customer, order, shipment, quantity, and date."),
]
table(d, ["Test", "Pass condition"], tests, [2.7, 4.4])

h1(d, "10. Ownership")
table(d, ["Owner", "Responsibility"], [
    ("Damon / admin", "Set product tracking requirements and approve policy changes."),
    ("Warehouse manager", "Configure stations, train operators, resolve blocked inventory, approve dispositions."),
    ("Warehouse operator", "Scan receive, pick, pack, and handoff. Cannot bypass required product fields."),
    ("Alex / engineering", "Implement product policies, ledger-safe workflows, scanner parsing, event notifications, and verification tests."),
    ("Finance", "Consume inventory valuation and reconciliation outputs. Does not bypass warehouse scan requirements."),
])

body(d, "Document location: Desktop/Unite_WMS_Scanner_Buy_and_Implementation_Guide.docx", color=GREY)

# Validate style constraints across paragraphs and table cells.
all_text = []
all_text.extend(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            all_text.extend(p.text for p in cell.paragraphs)
joined = "\n".join(all_text)
for banned in ("\u2014", "\u2013", ";"):
    if banned in joined:
        raise SystemExit(f"Banned punctuation found: {repr(banned)}")

OUT.parent.mkdir(parents=True, exist_ok=True)
d.save(OUT)
print(f"WROTE {OUT}")
